"""Cruza a aba "Capacidade ALE" (tancagem/modelo por base) com o Take or
Pay (TOP) apurado em Despesas.xlsx, e gera um CSV único por base+produto
pronto para o dashboard de capacidades.

Fontes:
- dados_tratados/capacidade_ale.csv (gerado por extract_capacidade_ale.py)
- arquivos_apoio/Despesas.xlsx (aba "Despesas")

Saída:
- dados_tratados/capacidade_ale_top.csv

Chave de junção: `unidade` (Capacidade ALE) normalizada (maiúsculas,
sem acento) == `FILIAL` (Despesas.xlsx) normalizada.

`codigo_base` ("Código Da Base") NÃO é uma chave confiável entre abas
diferentes do mesmo arquivo -- confirmado batendo errado para 3 bases
(Paulínia: 12 na Capacidade ALE vs 67 no Despesas.xlsx/Detalhamento
Tancagem; Rondonópolis: 69 vs 61; Cuiabá: 28 vs 73). O nome da unidade,
depois de normalizado, bate 100% em todas as fontes.

Definição de "tem TOP": a base tem pelo menos um contrato de
"Operação Normal" com Volume Mínimo (garantia contratual) preenchido e
maior que zero -- isso é o atributo contratual (existe cláusula Take or
Pay), independente de ela ter sido acionada (gerado pagamento) em algum
mês. O valor acumulado de TOP pago (`top_valor_realizado`) soma a coluna
"Take or Pay" apenas nas linhas com valor realizado (mesma regra usada em
dashboard_cessao_aggregate.py::_row_top), refletindo o quanto essa
cláusula já custou de fato.
"""

import csv
import json
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

BASE_DIR = Path(__file__).resolve().parent.parent
APOIO_DIR = BASE_DIR / "arquivos_apoio"
TRATADOS_DIR = BASE_DIR / "dados_tratados"

sys.path.insert(0, str(BASE_DIR / "scripts"))
from normalize import normalize_municipio  # noqa: E402

CAPACIDADE_CSV = TRATADOS_DIR / "capacidade_ale.csv"
DESPESAS_XLSX = APOIO_DIR / "Despesas.xlsx"
OUT_CSV = TRATADOS_DIR / "capacidade_ale_top.csv"


def _load_top_por_base():
    wb = openpyxl.load_workbook(DESPESAS_XLSX, data_only=True, read_only=True)
    ws = wb["Despesas"]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    idx = {h: i for i, h in enumerate(header) if h}

    congeneres_top = defaultdict(set)
    top_valor_realizado = defaultdict(float)
    meses_com_top = defaultdict(set)
    top_por_mes = defaultdict(lambda: defaultdict(float))

    for row in ws.iter_rows(min_row=2, values_only=True):
        filial = row[idx["FILIAL"]]
        if not filial:
            continue
        chave = normalize_municipio(filial)

        tipo = row[idx["TIPO DE COBRANÇA"]]
        vol_minimo = row[idx["Volume Mínimo "]] or 0
        if tipo == "Operação Normal" and vol_minimo > 0:
            congeneres_top[chave].add(row[idx["CONGENERE"]])

        valor_total = row[idx["Valor Total R$"]]
        if valor_total not in (None, 0):
            top_valor = row[idx["Take or Pay"]] or 0
            if top_valor:
                top_valor_realizado[chave] += top_valor
                data = row[idx["DATA CONTABIL"]]
                if data:
                    meses_com_top[chave].add((data.year, data.month))
                    period = f"{data.year}-{data.month:02d}"
                    top_por_mes[chave][period] += top_valor

    chaves = set(congeneres_top) | set(top_valor_realizado)
    return {
        chave: {
            "tem_top": chave in congeneres_top,
            "congeneres_top": "; ".join(sorted(congeneres_top.get(chave, []))),
            "top_valor_realizado": round(top_valor_realizado.get(chave, 0.0), 2),
            "meses_com_top_realizado": len(meses_com_top.get(chave, [])),
            "top_por_mes_json": json.dumps(
                {p: round(v, 2) for p, v in top_por_mes.get(chave, {}).items()}
            ),
        }
        for chave in chaves
    }


def build_capacidade_top():
    top_por_base = _load_top_por_base()

    with CAPACIDADE_CSV.open(encoding="utf-8") as f:
        rows = list(csv.DictReader(f))

    fieldnames = list(rows[0].keys()) + [
        "tem_top", "congeneres_top", "top_valor_realizado", "meses_com_top_realizado",
        "top_por_mes_json",
    ]
    out_rows = []
    for r in rows:
        chave = normalize_municipio(r["unidade"])
        info = top_por_base.get(chave, {
            "tem_top": False, "congeneres_top": "",
            "top_valor_realizado": 0.0, "meses_com_top_realizado": 0,
            "top_por_mes_json": "{}",
        })
        out_rows.append({**r, **info})

    OUT_CSV.parent.mkdir(exist_ok=True)
    with OUT_CSV.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(out_rows)

    return len(out_rows), OUT_CSV


if __name__ == "__main__":
    n, path = build_capacidade_top()
    print(f"{n} linhas escritas em {path}")
