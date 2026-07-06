"""Leitura da base XLSX e estruturas de dados da rodada.

Abas obrigatórias: 'Tabela de frete', 'Disponibilidade', 'Demanda'.
Abas opcionais:    'Estoque', 'Armazenagem'.

Mesmo schema do protótipo HTML, com uma diferença central: o MODAL de cada
rota é preservado (uma rota-modal por linha), em vez de colapsar tudo no
frete mais barato do par origem→destino.
"""

from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from .products import MODAL_UNKNOWN, loc_base, norm, norm_modal, norm_product

REQUIRED_SHEETS = ["Tabela de frete", "Disponibilidade", "Demanda"]


@dataclass
class Lane:
    """Rota-modal: frete R$/m³ de uma origem a um destino por um modal."""

    ko: str          # chave origem (normalizada)
    kd: str          # chave destino (normalizada)
    modal: str       # modal canônico (RODOVIARIO, FERROVIARIO, ...)
    modal_raw: str   # como veio na planilha (para diagnóstico)
    frete: float


@dataclass
class Supply:
    forn: str
    loc: str
    prod: str
    preco: float
    max: float
    med: float
    ko: str
    sid: str
    is_estoque: bool = False


@dataclass
class Demand:
    fil: str
    prod: str
    d: float
    kf: str


@dataclass
class Base:
    """Base da rodada, já tipada e normalizada."""

    lanes: list[Lane] = field(default_factory=list)
    supply: list[Supply] = field(default_factory=list)
    demand: list[Demand] = field(default_factory=list)
    arm: dict[str, float] = field(default_factory=dict)  # 'kf|PROD' ou 'kf|*' → R$/m³
    prods: list[str] = field(default_factory=list)
    blank_freight: int = 0
    est_n: int = 0
    est_vol: float = 0.0
    arm_n: int = 0
    src_name: str = ""

    def arm_cost(self, kf: str, prod: str) -> float:
        return self.arm.get(f"{kf}|{prod}", self.arm.get(f"{kf}|*", 0.0))

    def lane_index(self) -> dict[tuple[str, str], list[Lane]]:
        """(ko, kd) → rotas-modal disponíveis, menor frete por modal."""
        best: dict[tuple[str, str, str], Lane] = {}
        for ln in self.lanes:
            k = (ln.ko, ln.kd, ln.modal)
            if k not in best or ln.frete < best[k].frete:
                best[k] = ln
        idx: dict[tuple[str, str], list[Lane]] = {}
        for ln in best.values():
            idx.setdefault((ln.ko, ln.kd), []).append(ln)
        return idx


def _rows(ws) -> list[dict]:
    """Linhas da aba como dicts cabeçalho→valor (cabeçalho da linha 1)."""
    it = ws.iter_rows(values_only=True)
    header = next(it, None)
    if not header:
        return []
    keys = [str(h).strip() if h is not None else "" for h in header]
    return [dict(zip(keys, row)) for row in it if any(v is not None for v in row)]


def _pick(row: dict, *names):
    """Valor da primeira coluna cujo nome bate (case-insensitive, trim)."""
    lowered = {k.strip().lower(): v for k, v in row.items()}
    for n in names:
        if n.lower() in lowered:
            return lowered[n.lower()]
    return None


def _num(v) -> float | None:
    if v is None:
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f if f == f else None  # NaN → None


def load_base(path: str) -> Base:
    """Lê a base XLSX. Levanta ValueError se faltarem abas obrigatórias."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    missing = [s for s in REQUIRED_SHEETS if s not in wb.sheetnames]
    if missing:
        raise ValueError(
            "Abas ausentes: " + ", ".join(missing) + ". Gere o template e confira o formato."
        )
    sheet = lambda s: _rows(wb[s]) if s in wb.sheetnames else []
    base = Base()

    # --- Tabela de frete (modal preservado) ---
    for row in sheet("Tabela de frete"):
        o, d = _pick(row, "Origem"), _pick(row, "Destino")
        f = _num(_pick(row, "Frete (R$/m³)", "Frete"))
        if o is None or d is None or f is None or f <= 0:
            base.blank_freight += 1
            continue
        modal_raw = _pick(row, "Modal")
        ko = norm(_pick(row, "Chave Origem")) or norm(o)
        kd = norm(_pick(row, "Chave Destino")) or norm(d)
        base.lanes.append(
            Lane(ko=ko, kd=kd, modal=norm_modal(modal_raw),
                 modal_raw=str(modal_raw or "").strip(), frete=f)
        )

    # --- Demanda primeiro: dela o app aprende a carteira de produtos ---
    for row in sheet("Demanda"):
        fil = _pick(row, "Desc Filial", "Filial")
        prod = norm_product(_pick(row, "Produto"))
        tot = _num(_pick(row, "Demanda (m³)", "Total")) or 0
        if fil is None or not prod or prod.replace(".", "").isdigit() or tot <= 0:
            continue
        kf = norm(_pick(row, "Chave Filial")) or norm(fil)
        base.demand.append(Demand(fil=str(fil), prod=prod, d=tot, kf=kf))
    base.prods = sorted({d.prod for d in base.demand})

    # --- Disponibilidade (produto/preço tolerantes a cabeçalho trocado) ---
    i = 0
    for row in sheet("Disponibilidade"):
        forn = str(_pick(row, "Fornecedor") or "").strip()
        loc = _pick(row, "Localidade")
        if not forn or loc is None:
            continue
        cand = [_pick(row, "Produto"), _pick(row, "Nome")]
        prod = next((p for p in (norm_product(c) for c in cand) if p in base.prods), None)
        preco = _num(_pick(row, "Preço (R$/m³)", "Preco (R$/m³)"))
        if preco is None:
            preco = next((n for n in (_num(c) for c in cand) if n is not None and n > 0), None)
        if not prod or preco is None:
            continue
        ko = norm(_pick(row, "Chave Origem")) or norm(loc_base(loc))
        base.supply.append(Supply(
            forn=forn, loc=str(loc), prod=prod, preco=preco,
            max=_num(_pick(row, "Máx (m³)", "Max (m³)")) or 0,
            med=_num(_pick(row, "Média (m³)", "Media (m³)")) or 0,
            ko=ko, sid=f"{forn}|{ko}|{prod}#{i}",
        ))
        i += 1

    # --- Estoque: produto já disponível = origem adicional ---
    for row in sheet("Estoque"):
        local = _pick(row, "Local", "Localidade", "Filial")
        prod = norm_product(_pick(row, "Produto"))
        vol = _num(_pick(row, "Volume disponível (m³)", "Volume disponivel (m³)", "Volume (m³)")) or 0
        if local is None or not prod or vol <= 0:
            continue
        custo = _num(_pick(row, "Custo (R$/m³)", "Custo")) or 0.0
        ko = norm(local)
        base.supply.append(Supply(
            forn="ESTOQUE", loc=str(local), prod=prod, preco=custo,
            max=vol, med=vol, ko=ko, sid=f"ESTOQUE|{ko}|{prod}#{i}", is_estoque=True,
        ))
        i += 1
        base.est_n += 1
        base.est_vol += vol

    # --- Armazenagem: custo por filial (+produto opcional) ---
    for row in sheet("Armazenagem"):
        fil = _pick(row, "Filial", "Desc Filial")
        c = _num(_pick(row, "Custo armazenagem (R$/m³)", "Custo de Armazenagem", "Custo armazenagem"))
        if fil is None or c is None or c <= 0:
            continue
        prod = norm_product(_pick(row, "Produto"))
        kf = norm(fil)
        key = f"{kf}|{prod}" if prod in base.prods else f"{kf}|*"
        base.arm[key] = c
        base.arm_n += 1

    wb.close()
    return base


TEMPLATE_SHEETS = {
    "Tabela de frete": ["Origem", "UF Origem", "Destino", "UF Destino", "Frete (R$/m³)", "Modal"],
    "Disponibilidade": ["Fornecedor", "Localidade", "Produto", "Preço (R$/m³)",
                        "Volume mínimo (m³)", "Média (m³)", "Máx (m³)"],
    "Demanda": ["Desc Filial", "Produto", "Demanda (m³)"],
    "Estoque": ["Local", "Produto", "Volume disponível (m³)", "Custo (R$/m³)"],
    "Armazenagem": ["Filial", "Produto", "Custo armazenagem (R$/m³)"],
}


def make_template(path: str) -> None:
    """Gera o template guiado: 5 abas com a geografia da malha pré-preenchida.

    As linhas vêm de template_data.json (portado do protótipo) — nomes de
    origens/destinos/filiais já reconciliados entre abas, valores em branco
    para preencher. Quem preenche digita números, não geografia.
    """
    import json

    data = json.loads((Path(__file__).parent / "template_data.json").read_text(encoding="utf-8"))
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, headers in TEMPLATE_SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in data.get(name, []):
            ws.append([row.get(h) for h in headers])
    wb.save(path)
