"""Integração ponta a ponta: XLSX entra → LP → XLSX sai (com auditoria)."""

import openpyxl
import pytest

from bussola.io_excel import TEMPLATE_SHEETS, load_base, make_template
from bussola.model import solve
from bussola.report import export_xlsx, modal_mix_df


def write_base(path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    data = {
        "Tabela de frete": [
            ["Origem", "UF Origem", "Destino", "UF Destino", "Frete (R$/m³)", "Modal"],
            ["X", "SP", "Fil A", "MG", 120, "Rodo"],
            ["X", "SP", "Fil A", "MG", 80, "Ferro"],
            ["Y", "RJ", "Fil A", "MG", 90, "Rodo"],
            ["X", "SP", "Fil B", "PR", 70, "Rodo"],
            ["X", "SP", None, None, None, None],  # linha em branco: ignorada
        ],
        "Disponibilidade": [
            ["Fornecedor", "Localidade", "Produto", "Preço (R$/m³)",
             "Volume mínimo (m³)", "Média (m³)", "Máx (m³)"],
            ["F1", "X", "Diesel S10", 3000, 0, 0, 80],   # alias → S10
            ["F2", "Y", "S10", 3100, 0, 0, 100],
            ["F1", "X", "S500", 2900, 0, 0, 60],
        ],
        "Demanda": [
            ["Desc Filial", "Produto", "Demanda (m³)"],
            ["Fil A", "S10", 100],
            ["Fil B", "S500", 50],
        ],
        "Estoque": [
            ["Local", "Produto", "Volume disponível (m³)", "Custo (R$/m³)"],
            ["Fil B", "S500", 10, 0],
        ],
        "Armazenagem": [
            ["Filial", "Produto", "Custo armazenagem (R$/m³)"],
        ],
    }
    for name, rows in data.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    wb.save(path)


def test_roundtrip(tmp_path):
    src = tmp_path / "base.xlsx"
    write_base(src)
    base = load_base(str(src))

    assert base.prods == ["S10", "S500"]
    assert base.blank_freight == 1
    assert base.est_n == 1 and base.est_vol == 10
    assert len(base.lanes) == 4

    res = solve(base)
    assert res.feasible and not res.ruptures
    # estoque local (R$ 0/m³) cobre 10 m³ de S500; X→B cobre os 40 restantes
    # S10: 80×3.080 + 20×3.190 = 310.200 · S500: 40×2.970 = 118.800
    assert res.total_cost == pytest.approx(310_200 + 118_800)

    mix = modal_mix_df(res)
    assert set(mix["Modal"]) == {"FERROVIARIO", "RODOVIARIO", "LOCAL"}

    out = tmp_path / "resultado.xlsx"
    export_xlsx(res, base, str(out), src_name="base.xlsx")
    wb = openpyxl.load_workbook(str(out))
    assert wb.sheetnames == ["Plano de atendimento", "Mix modal", "Consumo por polo",
                             "Atendimento da demanda", "KPIs", "Auditoria"]
    plano = wb["Plano de atendimento"]
    header = [c.value for c in plano[1]]
    assert "Modal" in header  # o modal usado é reportado por fluxo


def test_missing_sheet_raises(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Demanda"
    p = tmp_path / "ruim.xlsx"
    wb.save(p)
    with pytest.raises(ValueError, match="Abas ausentes"):
        load_base(str(p))


def test_template_has_all_sheets(tmp_path):
    p = tmp_path / "template.xlsx"
    make_template(str(p))
    wb = openpyxl.load_workbook(str(p))
    assert wb.sheetnames == list(TEMPLATE_SHEETS)
    assert [c.value for c in wb["Tabela de frete"][1]] == TEMPLATE_SHEETS["Tabela de frete"]
